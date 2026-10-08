#!/usr/bin/env python3
"""Export StringID / English / Korean reference pairs for translators.

Matches the Korean patch against a ROM (default: vanilla Explorers of Alpha)
and writes only strings that received a Korean translation.

Usage: python export_reference.py [rom.nds] [out.xlsx]
"""

from __future__ import annotations

import sys
from pathlib import Path

from ndspy.rom import NintendoDSRom
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from skytemple_files.common.util import get_ppmdu_config_for_rom

MODULE_DIR = Path(__file__).resolve().parent.parent
REPO = MODULE_DIR.parent.parent
sys.path.insert(0, str(REPO))

from patch_engine.apply_korean import (  # noqa: E402
    TRANSLATED_KINDS,
    load_source,
    plan_translation,
)
from patch_engine.manifest import load_module_manifest  # noqa: E402

GLOSSARY_BLOCKS = [
    "Pokemon Names",
    "New Pokemon Names",
    "Move Names",
    "Item Names",
    "Ability Names",
    "Type Names",
    "Status Names and Descriptions",
    "IQ Skills Names",
    "Tactics Names",
    "Trap Names",
    "Weather Names",
    "Explorer Ranks Names",
    "Dungeon Names (Main)",
    "Ground Map Names",
    "Pokemon Categories",
]


def _cell(value):
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


def _sheet(wb: Workbook, title: str, header: list[str], widths: list[int], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for row in rows:
        ws.append([_cell(v) for v in row])
    wrap = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main() -> None:
    rom_path = Path(sys.argv[1]) if len(sys.argv) > 1 else REPO / "PatchTesting" / "Explorers of Alpha" / "Explorers of Alpha.nds"
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else REPO / "PatchTesting" / "Export Rom" / "Explorers of Alpha_Vanilla_korean_reference.xlsx"

    rom = NintendoDSRom(rom_path.read_bytes())
    blocks = sorted(
        get_ppmdu_config_for_rom(rom).string_index_data.string_blocks.items(),
        key=lambda kv: kv[1].begin,
    )

    def block_of(sid: int) -> str:
        for name, b in blocks:
            if b.begin <= sid < b.end:
                return name
        return ""

    src = load_source(MODULE_DIR, load_module_manifest(MODULE_DIR))
    plan = plan_translation(rom, src)

    text_rows = [
        [j, plan.text_e[j], d.ko, block_of(j)]
        for j, d in enumerate(plan.text_dec)
        if d.kind in TRANSLATED_KINDS
    ]
    script_rows = [
        [f"{path}#{k}", plan.scripts[path][k], d.ko]
        for path, decs in plan.script_dec.items()
        for k, d in enumerate(decs)
        if d.kind in TRANSLATED_KINDS
    ]
    by_block = {name: b for name, b in blocks}
    glossary: list[list] = []
    for name in GLOSSARY_BLOCKS:
        b = by_block.get(name)
        if b is None:
            continue
        seen: set[str] = set()
        for j in range(b.begin, min(b.end, len(plan.text_dec))):
            d = plan.text_dec[j]
            en = plan.text_e[j]
            if d.kind in TRANSLATED_KINDS and en not in seen:
                seen.add(en)
                glossary.append([name, j, en, d.ko])

    wb = Workbook()
    ws = wb.active
    ws.title = "안내"
    for line in (
        ["Explorers of Alpha 한글 번역 참고 자료"],
        [f"기준 ROM: {rom_path.name}"],
        ["한글: 비공식 EoS 한글패치(하늘의 탐험대 1-101b) 번역 중, 영어 원문이 EoS 원문과 일치하는 문장만 수록"],
        ["StringID: text_e.str 문자열 번호 / 스크립트는 '파일경로#문자열 번호'"],
        ["태그 규칙: [CS:x]…[CR] 색상, [LS:n]…[LE] 용어 링크, [K] 입력 대기, [C] 페이지 넘김, [CN] 가운데 정렬,"],
        ["          [string:n]/[value:n:m]/[kind:n] 등은 게임이 채우는 변수입니다. 태그는 그대로 유지합니다."],
        ["줄바꿈(\\n)은 대사창 한 줄을 뜻합니다. 한글은 영어보다 한 줄에 들어가는 글자 수가 적습니다."],
        [f"text_e {len(text_rows)}개 / 스크립트 {len(script_rows)}개 / 용어집 {len(glossary)}개"],
    ):
        ws.append(line)
    ws.column_dimensions["A"].width = 120

    _sheet(wb, "text_e", ["StringID", "영어 원문", "한글", "분류"], [9, 60, 60, 30], text_rows)
    _sheet(wb, "스크립트", ["StringID", "영어 원문", "한글"], [36, 60, 60], script_rows)
    _sheet(wb, "용어집", ["분류", "StringID", "영어", "한글"], [28, 9, 30, 30], glossary)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"text_e {len(text_rows)}, scripts {len(script_rows)}, glossary {len(glossary)} -> {out}")


if __name__ == "__main__":
    main()
