#!/usr/bin/env python3
"""Import a translated copy of the untranslated-strings report into data/translations.json.gz.

Reads the "text_e 미번역" and "스크립트 미번역" sheets. The "한국어 번역" column is the
translation; the English column is kept so apply can skip strings whose Alpha
English has changed since. Empty translations are skipped.

Usage: python import_translations.py <translated.xlsx> [out.json.gz]
"""

from __future__ import annotations

import gzip
import json
import re
import sys
from pathlib import Path

import openpyxl

MODULE_DIR = Path(__file__).resolve().parent.parent
TEXT_SHEET = "text_e 미번역"
SCRIPT_SHEET = "스크립트 미번역"
KO_HEADER = "한국어 번역"
_XML_CR = re.compile(r"_x000[Dd]_")


def _clean(value) -> str:
    if value is None:
        return ""
    s = _XML_CR.sub("", str(value))
    return s.replace("\r\n", "\n").replace("\r", "\n")


def _columns(header: tuple, *names: str) -> list[int]:
    cols = [str(h).strip() if h is not None else "" for h in header]
    missing = [n for n in names if n not in cols]
    if missing:
        raise SystemExit(f"columns {missing} not found in header {cols}")
    return [cols.index(n) for n in names]


def read_translations(xlsx: Path) -> dict:
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    text_rows = list(wb[TEXT_SHEET].iter_rows(values_only=True))
    c_id, c_en, c_ko = _columns(text_rows[0], "String ID", "Alpha 영어 (현재)", KO_HEADER)
    text_e: dict[str, dict[str, str]] = {}
    for r in text_rows[1:]:
        ko = _clean(r[c_ko])
        if r[c_id] is None or not ko:
            continue
        text_e[str(int(r[c_id]))] = {"en": _clean(r[c_en]), "ko": ko}

    script_rows = list(wb[SCRIPT_SHEET].iter_rows(values_only=True))
    c_file, c_idx, c_en, c_ko = _columns(script_rows[0], "파일", "번호", "Alpha 영어 (현재)", KO_HEADER)
    scripts: dict[str, dict[str, dict[str, str]]] = {}
    for r in script_rows[1:]:
        ko = _clean(r[c_ko])
        if r[c_file] is None or not ko:
            continue
        scripts.setdefault(str(r[c_file]), {})[str(int(r[c_idx]))] = {"en": _clean(r[c_en]), "ko": ko}
    return {"source": xlsx.name, "text_e": text_e, "scripts": scripts}


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    xlsx = Path(sys.argv[1])
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else MODULE_DIR / "data" / "translations.json.gz"
    doc = read_translations(xlsx)
    raw = json.dumps(doc, ensure_ascii=False, sort_keys=True, indent=0).encode("utf-8")
    with gzip.GzipFile(out, "wb", mtime=0) as fh:
        fh.write(raw)
    n_scripts = sum(len(v) for v in doc["scripts"].values())
    print(f"{out}: text_e {len(doc['text_e'])}, scripts {n_scripts} in {len(doc['scripts'])} files")


if __name__ == "__main__":
    main()
