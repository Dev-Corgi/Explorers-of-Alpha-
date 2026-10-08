"""Excel report of strings the Korean module leaves in English."""

from __future__ import annotations

import difflib
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

from .apply_korean import (
    GLOBAL,
    ICON_PREFIX,
    MOVED,
    SAME_FILE,
    SOURCE_ENGLISH,
    UNMATCHED,
    USER,
    KoreanPlan,
    KoreanSource,
    plan_stats,
)

KIND_LABEL = {
    "same_id": "같은 ID/번호 영어 일치",
    MOVED: "ID 이동 (같은 영어, 다른 ID)",
    ICON_PREFIX: "아이콘 태그만 다름 (본문 영어 일치)",
    SAME_FILE: "같은 파일 다른 번호 (같은 영어)",
    GLOBAL: "다른 스크립트의 같은 영어",
    UNMATCHED: "대응 영어 없음 (Alpha 수정/추가 문장)",
    SOURCE_ENGLISH: "원본 한글패치도 영어 유지",
    USER: "번역 엑셀 적용 (data/translations.json.gz)",
    "no_text": "번역 대상 아님 (태그/빈 문자열)",
}
SUGGEST_WINDOW_TEXT = 60
SUGGEST_WINDOW_SCRIPT = 30
SUGGEST_MIN = 0.6


def _cell(value):
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)[:32000]
    return value


def _best(target: str, cands: list[tuple[object, str]]) -> tuple[object, float] | None:
    sm = difflib.SequenceMatcher(autojunk=False)
    sm.set_seq2(target)
    best: tuple[object, float] | None = None
    for ref, text in cands:
        if not text:
            continue
        sm.set_seq1(text)
        floor = best[1] if best else SUGGEST_MIN
        if sm.real_quick_ratio() < floor or sm.quick_ratio() < floor:
            continue
        r = sm.ratio()
        if r >= floor:
            best = (ref, r)
    return best


def _sheet(wb: Workbook, title: str, header: list[str], widths: list[int], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append([_cell(v) for v in row])
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_korean_report(plan: KoreanPlan, src: KoreanSource, path: Path) -> dict[str, int]:
    en, ko = src.text_en, src.text_ko

    text_rows: list[list] = []
    for j, d in enumerate(plan.text_dec):
        if d.kind not in (UNMATCHED, SOURCE_ENGLISH):
            continue
        ref_id, ref_en, ref_ko, ratio = "", "", "", ""
        if d.kind == UNMATCHED:
            exp = plan.expected_eos_id[j]
            lo, hi = max(0, exp - SUGGEST_WINDOW_TEXT), min(len(en), exp + SUGGEST_WINDOW_TEXT + 1)
            best = _best(plan.text_e[j], [(i, en[i]) for i in range(lo, hi)])
            if best:
                i = int(best[0])
                ref_id, ref_en, ref_ko, ratio = i, en[i], ko[i], round(best[1], 2)
        text_rows.append([j, KIND_LABEL[d.kind], plan.text_e[j], ref_id, ref_en, ref_ko, ratio])

    script_rows: list[list] = []
    for p, decs in plan.script_dec.items():
        s = src.scripts.get(p)
        for k, d in enumerate(decs):
            if d.kind not in (UNMATCHED, SOURCE_ENGLISH):
                continue
            ref_en, ref_ko, ratio = "", "", ""
            if d.kind == UNMATCHED and s:
                lo, hi = max(0, k - SUGGEST_WINDOW_SCRIPT), min(len(s["en"]), k + SUGGEST_WINDOW_SCRIPT + 1)
                best = _best(plan.scripts[p][k], [(i, s["en"][i]) for i in range(lo, hi)])
                if best:
                    i = int(best[0])
                    ref_en, ref_ko, ratio = s["en"][i], s["ko"][i], round(best[1], 2)
            note = "" if s else "Alpha 전용 스크립트"
            script_rows.append([p, k, KIND_LABEL[d.kind], plan.scripts[p][k], ref_en, ref_ko, ratio, note])

    moved_rows = [
        [j, d.ref, KIND_LABEL[d.kind], plan.text_e[j], d.ko]
        for j, d in enumerate(plan.text_dec)
        if d.kind in (MOVED, ICON_PREFIX)
    ]
    cross_rows = [
        [p, k, KIND_LABEL[d.kind], d.ref, plan.scripts[p][k], d.ko]
        for p, decs in plan.script_dec.items()
        for k, d in enumerate(decs)
        if d.kind in (SAME_FILE, GLOBAL)
    ]

    stats = plan_stats(plan)
    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    ws.append(["구분", "분류", "문자열 수"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for section, counts in (("text_e.str", stats["text_e"]), ("스크립트(.ssb)", stats["scripts"])):
        for kind, label in KIND_LABEL.items():
            if kind in counts:
                ws.append([section, label, counts[kind]])
    ws.append([])
    ws.append(["적용 기준", "Alpha 영어가 한글패치 원본(EoS) 영어와 정확히 같을 때만 그 한글을 적용합니다."])
    ws.append(["", "아이콘 태그만 다른 문장은 Alpha 태그를 유지하고 본문만 한글로 바꿉니다."])
    ws.append(["", "미번역 시트의 '참고' 열은 비슷한 EoS 문장과 그 한글(자동 추정)이며 적용되지 않았습니다."])
    ws.append(["", "번역 엑셀 문장은 그 영어가 지금 롬의 영어와 같을 때만 적용하고, 다르면 '검토_번역 불일치'에 남깁니다."])
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 44
    ws.column_dimensions["C"].width = 12

    _sheet(
        wb, "text_e 미번역",
        ["String ID", "분류", "Alpha 영어 (현재)", "참고 EoS ID", "참고 EoS 영어", "참고 한글", "유사도"],
        [10, 26, 60, 11, 50, 50, 8], text_rows,
    )
    _sheet(
        wb, "스크립트 미번역",
        ["파일", "번호", "분류", "Alpha 영어 (현재)", "참고 EoS 영어", "참고 한글", "유사도", "비고"],
        [30, 7, 26, 60, 50, 50, 8, 16], script_rows,
    )
    _sheet(
        wb, "검토_text_e ID이동",
        ["Alpha ID", "EoS ID", "방식", "Alpha 영어", "적용한 한글"],
        [10, 10, 30, 50, 50], moved_rows,
    )
    _sheet(
        wb, "검토_스크립트 교차",
        ["파일", "번호", "방식", "한글 출처", "Alpha 영어", "적용한 한글"],
        [30, 7, 26, 36, 50, 50], cross_rows,
    )
    stale_rows = [
        [where, i, en, _current(plan, where, i), ko] for where, i, en, ko in plan.user_stale
    ]
    _sheet(
        wb, "검토_번역 불일치",
        ["파일/text_e", "번호", "번역 당시 영어", "지금 롬 영어", "번역 (미적용)"],
        [30, 7, 50, 50, 50], stale_rows,
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {
        "text_e_untranslated": len(text_rows),
        "script_untranslated": len(script_rows),
        "text_e_moved": len(moved_rows),
        "script_cross": len(cross_rows),
        "translations_stale": len(stale_rows),
    }


def _current(plan: KoreanPlan, where: str, i: int) -> str:
    strings = plan.text_e if where == "text_e" else plan.scripts.get(where, [])
    return strings[i] if i < len(strings) else "(없음)"
